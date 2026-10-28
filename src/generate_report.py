from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from data_processor import HousePricesDataProcessor, PreprocessingReport
from model_trainer import HousePricesModelTrainer, TrainingOutcome
from visualizer import HousePricesVisualizer


def build_markdown_report(
    output_md: Path,
    prep: PreprocessingReport,
    outcome: TrainingOutcome,
    figures_rel_dir: str = "../outputs/figures",
) -> None:
    top3 = prep.top_correlations.head(3)
    top10 = prep.top_correlations.head(10)

    r2_base = outcome.baseline_metrics["r2"]
    mae_base = outcome.baseline_metrics["mae"]
    rmse_base = outcome.baseline_metrics["rmse"]
    mape_base = outcome.baseline_metrics["mape"]

    r2_comp = outcome.competition_metrics["r2"]
    mae_comp = outcome.competition_metrics["mae"]
    rmse_comp = outcome.competition_metrics["rmse"]
    mape_comp = outcome.competition_metrics["mape"]

    corr_rows = "\n".join(
        [f"| {idx + 1}º | `{col}` | **{val:.4f}** |" for idx, (col, val) in enumerate(top10.items())]
    )

    null_rows = "\n".join(
        [
            f"| `{row['Coluna']}` | `{row['Tipo']}` | {row['Qtd_Nulos']} | {row['Percentual (%)']:.2f}% |"
            for _, row in prep.null_summary.head(10).iterrows()
        ]
    )

    content = f"""# Relatório Técnico — Trabalho Prático AV1 (Machine Learning)
## Parte 1: Aprendizado Supervisionado — Problema 1: Regressão Linear ("O Preço Justo")

**Base de Dados Analisada:** `house-prices.csv` ({prep.total_rows} registros × {prep.total_cols} variáveis)  
**Variável Alvo (Target):** `SalePrice` (Preço de venda do imóvel em US$)

---

## 1. Análise Exploratória de Dados (EDA)

> **Pergunta:** *Quais são as 3 variáveis que possuem a maior correlação linear positiva com o preço de venda (`SalePrice`)? Mostre isso através de um gráfico de dispersão (scatter plot) ou matriz de correlação.*

### Resposta Direta
Calculando o Coeficiente de Correlação Linear de Pearson ($r$) entre todas as variáveis numéricas do dataset e a variável alvo `SalePrice`, as **3 variáveis com maior correlação linear positiva** são:

1. **`OverallQual` ($r = {top3.iloc[0]:.4f}$):** Avalia a qualidade geral do material e do acabamento da construção (escala de 1 a 10). É o principal fator determinante de valorização imobiliária na base.
2. **`GrLivArea` ($r = {top3.iloc[1]:.4f}$):** Área útil habitável acima do nível do solo (em pés quadrados / *square feet*). Quanto maior a metragem construída útil, maior o valor de mercado.
3. **`GarageCars` ($r = {top3.iloc[2]:.4f}$):** Capacidade da garagem em número de automóveis (seguida de perto por `GarageArea` com $r = {top10.iloc[3]:.4f}$ e `TotalBsmtSF` com $r = {top10.iloc[4]:.4f}$), refletindo tanto o padrão construtivo do terreno quanto a infraestrutura do imóvel.

### Ranking das 10 Maiores Correlações Positivas com `SalePrice`

| Posição | Variável | Correlação de Pearson ($r$) |
| :---: | :--- | :---: |
{corr_rows}

### Evidências Visuais (Matriz de Correlação e Gráficos de Dispersão)

![Matriz de Correlação]({figures_rel_dir}/01_matriz_correlacao.png)

![Gráficos de Dispersão Top 3]({figures_rel_dir}/02_scatter_top3_correlacao.png)

---

## 2. Tratamento de Dados (Missing Values)

> **Pergunta:** *O dataset possui valores nulos (missing values). Escolha duas colunas com valores nulos: uma numérica e uma categórica. Como você tratou cada uma delas (ex: preencheu com a média, moda, ou apagou a linha)? Justifique sua escolha.*

### Diagnóstico de Nulos no Dataset
Das {prep.total_cols} colunas originais, **19 apresentam valores nulos**. Abaixo estão as 10 variáveis com maior incidência de ausências:

| Coluna | Tipo de Dado | Qtd. Nulos | % sobre o Total ({prep.total_rows} linhas) |
| :--- | :--- | :---: | :---: |
{null_rows}

### Tratamento Escolhido e Justificativa Técnica

1. **Variável Numérica Escolhida — `LotFrontage` ({prep.numeric_null_count} valores nulos / 17.80% da base):**
   - **Tratamento aplicado:** Imputação pela **mediana** (`{prep.numeric_fill_value:.1f}` pés lineares).
   - **Justificativa técnica:** A exclusão das 259 linhas (`dropna`) descartaria **17.8% da base de treino**, reduzindo drasticamente o poder estatístico do modelo e inviabilizando a predição caso algum dos 5 imóveis secretos da competição possua `LotFrontage` nulo. Além disso, como a distribuição de testada de lotes possui assimetria positiva (*outliers* com terrenos de mais de 300 pés), a **mediana** (`69.0` ft) é um estimador de tendência central mais robusto que a média (`70.0` ft), evitando distorcer os coeficientes da regressão.

2. **Variável Categórica Escolhida — `GarageType` ({prep.categorical_null_count} valores nulos / 5.43% da base):**
   - **Tratamento aplicado:** Preenchimento semântico com uma nova categoria constante **`"None"`** (Sem Garagem).
   - **Justificativa técnica:** Na documentação oficial da base *Ames Housing*, valores `NA` em atributos de infraestrutura (`GarageType`, `FireplaceQu`, `PoolQC`, `Alley`, `Fence`, `BsmtQual`) **não representam erro de coleta**, mas sim a **ausência física da estrutura no imóvel** (imóveis sem garagem possuem `GarageCars = 0` e `GarageArea = 0`). Se imputássemos pela moda (`"Attchd"` — garagem anexada à casa), introduziríamos um erro grave de negócio ao atribuir uma garagem inexistente a 79 imóveis. Criar a categoria `"None"` preserva 100% da integridade semântica dos dados.

*(Nota: o mesmo critério técnico foi estendido às demais variáveis da base para garantir a matriz completa no treinamento: `"None"` nas categóricas de ausência física, moda em `Electrical` que possui apenas 1 nulo, e mediana em `MasVnrArea` e `GarageYrBlt`).*

![Tratamento de Valores Nulos]({figures_rel_dir}/03_tratamento_valores_nulos.png)

---

## 3. Engenharia de Atributos (Codificação Categórica)

> **Pergunta:** *Transforme pelo menos duas variáveis categóricas em numéricas (usando One-Hot Encoding ou Label Encoding). Explique por que o modelo de Regressão Linear exige essa conversão.*

### Transformações Realizadas

Aplicamos as duas técnicas de acordo com a natureza matemática de cada variável categórica:

1. **Ordinal / Label Encoding em `ExterQual` e `KitchenQual` (e demais variáveis de qualidade/condição):**
   - Essas variáveis possuem uma **ordem hierárquica intrínseca** de qualidade: `None` (0) < `Po` (1 - *Poor*) < `Fa` (2 - *Fair*) < `TA` (3 - *Typical/Average*) < `Gd` (4 - *Good*) < `Ex` (5 - *Excellent*).
   - Convertê-las para inteiros de `0` a `5` preserva a relação monotônica direta com o preço do imóvel em uma única coluna numérica.

2. **One-Hot Encoding (`pd.get_dummies(..., drop_first=True)`) em `Neighborhood`, `GarageType` e `MSZoning`:**
   - São variáveis **nominais puras** (não existe hierarquia natural entre bairros ou tipos de garagem).
   - Foram expandidas em variáveis binárias (*dummy variables* `0` ou `1`), utilizando `drop_first=True` ($k - 1$ colunas) para evitar a **armadilha da multicolinearidade perfeita** (*Dummy Variable Trap*) na inversão da matriz $(X^T X)^{-1}$.
3. **Feature Engineering de Domínio Imobiliário (Modelo Hedônico):**
   - Criamos variáveis compostas de alto poder explicativo linear: Área Construída Total (`TotalSF = TotalBsmtSF + 1stFlrSF + 2ndFlrSF`), Total Ponderado de Banheiros (`TotalBath`), Idade do Imóvel (`HouseAge`), interações qualidade × metragem (`Qual_x_TotalSF = OverallQual * TotalSF`) e Índice Hedônico de Preço Mediano por pé quadrado do Bairro (`Neigh_PPSF`, estimado exclusivamente no conjunto de treino para evitar *data leakage*).
   - Após a codificação completa, o espaço de atributos passou de **{prep.features_before_encoding}** colunas originais para **{prep.features_after_encoding}** features numéricas.

### Por que o Modelo de Regressão Linear Exige Essa Conversão?
A Regressão Linear Múltipla modela a variável resposta como uma combinação linear algébrica dos atributos de entrada:

$$\\hat{{y}} = \\beta_0 + \\beta_1 x_1 + \\beta_2 x_2 + \\dots + \\beta_p x_p$$

Para estimar o vetor de pesos $\\hat{{\\beta}} = (X^T X)^{{-1}} X^T y$ via Mínimos Quadrados Ordinários (OLS), o algoritmo realiza multiplicações de matrizes, produtos escalares e cálculo de gradientes em um espaço vetorial real ($\\mathbb{{R}}^n$). Strings ou textos como `"CollgCr"` ou `"Attchd"` não possuem representação aritmética nem distância métrica definida; portanto, precisam ser mapeados para valores numéricos para que o hiperplano de regressão possa ser calculado.

---

## 4. Treinamento, Baseline e Interpretação de Negócios

> **Pergunta:** *Separe os dados em treino e teste (ex: 80/20). Treine um modelo de Regressão Linear múltipla básico. Sem fazer ajustes avançados, qual foi o R² (R-quadrado) obtido nos dados de teste? O que esse número significa para a imobiliária?*

### Configuração Experimental e Resultados
Os dados pré-processados ({prep.total_rows} imóveis) foram divididos em:
- **Conjunto de Treino (80%):** `{outcome.train_size}` registros (`random_state=42`)
- **Conjunto de Teste (20%):** `{outcome.test_size}` registros

Além do **Modelo Baseline (`LinearRegression` puro, sem ajustes avançados)** exigido pelo enunciado, treinamos também um **Modelo Otimizado para a Competição** (Ensemble combinando `Ridge`, `Lasso` e `GradientBoostingRegressor` com transformação logarítmica $\\log(1 + \\text{{SalePrice}})$ focado na minimização direta do MAE para os 5 imóveis secretos):

| Modelo | $R^2$ (Teste) | MAE — Erro Absoluto Médio (US$) | RMSE (US$) | MAPE (%) |
| :--- | :---: | :---: | :---: | :---: |
| **1. Regressão Linear Múltipla (Baseline Puro)** | **{r2_base:.4f} ({r2_base * 100:.2f}%)** | US$ {mae_base:,.2f} | US$ {rmse_base:,.2f} | {mape_base:.2f}% |
| **2. Pipeline Otimizado (Competição — 5 Imóveis)** | **{r2_comp:.4f} ({r2_comp * 100:.2f}%)** | **US$ {mae_comp:,.2f}** | US$ {rmse_comp:,.2f} | {mape_comp:.2f}% |

![Avaliação de Modelos]({figures_rel_dir}/04_avaliacao_regressao_linear.png)

### O que o $R^2 = {r2_base:.4f}$ significa para a Imobiliária?
1. **Explicação da Variância de Preços:** O Coeficiente de Determinação ($R^2 = {r2_base:.4f}$) indica que **{r2_base * 100:.2f}% de toda a variação nos preços de venda dos imóveis** da carteira de teste é explicada diretamente pelas características físicas e de localização capturadas pelo modelo de Regressão Linear Múltipla.
2. **Confiabilidade na Precificação Automática:** Para a operação da imobiliária, um $R^2$ acima de 85% já no baseline demonstra que o modelo possui forte aderência ao mercado imobiliário local, reduzindo a subjetividade humana de corretores ao avaliar novos imóveis captados.
3. **Margem Financeira e Risco Operacional (MAE):** Na prática comercial, o modelo baseline erra em média **US$ {mae_base:,.2f}** (MAE) para mais ou para menos em relação ao valor real fechado. Já com o **Pipeline Otimizado para a Competição**, o $R^2$ sobe para **{r2_comp:.4f} ({r2_comp * 100:.2f}%)** e o Erro Absoluto Médio cai para **US$ {mae_comp:,.2f}** ({mape_comp:.2f}% de erro percentual médio), oferecendo uma precificação de alta precisão para vencer o desafio dos 5 imóveis secretos.
"""
    output_md.write_text(content, encoding="utf-8")


def build_docx_report(
    output_docx: Path,
    prep: PreprocessingReport,
    outcome: TrainingOutcome,
    chart_paths: dict,
) -> None:
    doc = Document()

    # Base style adjustments
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)

    title = doc.add_heading("Relatório Técnico — Trabalho Prático AV1 (Machine Learning)", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = subtitle.add_run(
        'Parte 1: Aprendizado Supervisionado — Problema 1: Regressão Linear ("O Preço Justo")'
    )
    run_sub.bold = True
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(30, 58, 138)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        f"Dataset: house-prices.csv ({prep.total_rows} registros × {prep.total_cols} colunas) | Variável Alvo: SalePrice (US$)"
    ).italic = True

    # ---------------- QUESTION 1 ----------------
    doc.add_heading("1. Análise Exploratória de Dados (EDA)", level=1)
    q1_p = doc.add_paragraph()
    q1_r = q1_p.add_run(
        "Pergunta: Quais são as 3 variáveis que possuem a maior correlação linear positiva com o preço de venda (SalePrice)? "
        "Mostre isso através de um gráfico de dispersão (scatter plot) ou matriz de correlação."
    )
    q1_r.italic = True

    top3 = prep.top_correlations.head(3)
    top10 = prep.top_correlations.head(10)

    doc.add_paragraph(
        "Por meio do Coeficiente de Correlação Linear de Pearson (r) calculado sobre todas as variáveis numéricas do "
        "dataset em relação ao preço de venda (SalePrice), identificamos que as 3 variáveis com maior correlação linear positiva são:"
    )

    b1 = doc.add_paragraph(style="List Number")
    b1.add_run(f"OverallQual (r = {top3.iloc[0]:.4f}): ").bold = True
    b1.add_run(
        "Avalia a qualidade geral dos materiais e do acabamento da residência (escala de 1 a 10). "
        "É o preditor isolado mais forte de valorização do imóvel."
    )

    b2 = doc.add_paragraph(style="List Number")
    b2.add_run(f"GrLivArea (r = {top3.iloc[1]:.4f}): ").bold = True
    b2.add_run(
        "Representa a área útil habitável acima do nível do solo (em pés quadrados / sq ft). "
        "Apresenta forte relação linear positiva com o preço final de venda."
    )

    b3 = doc.add_paragraph(style="List Number")
    b3.add_run(f"GarageCars (r = {top3.iloc[2]:.4f}): ").bold = True
    b3.add_run(
        f"Capacidade da garagem medida pelo número de vagas de veículos (seguida por GarageArea com r = {top10.iloc[3]:.4f} "
        f"e TotalBsmtSF com r = {top10.iloc[4]:.4f}), atuando como indicador direto do porte construtivo do imóvel."
    )

    # Table of Top 8 Correlations
    table_corr = doc.add_table(rows=1, cols=3)
    table_corr.style = "Light Shading Accent 1"
    table_corr.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table_corr.rows[0].cells
    hdr[0].text = "Posição"
    hdr[1].text = "Variável"
    hdr[2].text = "Correlação de Pearson (r)"
    for idx, (col, val) in enumerate(top10.head(8).items()):
        row_cells = table_corr.add_row().cells
        row_cells[0].text = f"{idx + 1}º"
        row_cells[1].text = col
        row_cells[2].text = f"{val:.4f}"

    doc.add_paragraph()
    doc.add_picture(str(chart_paths["corr"]), width=Inches(5.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_picture(str(chart_paths["scatter"]), width=Inches(6.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ---------------- QUESTION 2 ----------------
    doc.add_heading("2. Tratamento de Dados (Valores Nulos)", level=1)
    q2_p = doc.add_paragraph()
    q2_p.add_run(
        "Pergunta: O dataset possui valores nulos (missing values). Escolha duas colunas com valores nulos: uma numérica "
        "e uma categórica. Como você tratou cada uma delas (ex: preencheu com a média, moda, ou apagou a linha)? Justifique sua escolha."
    ).italic = True

    doc.add_paragraph(
        f"A inspeção estrutural do dataset revelou que 19 das {prep.total_cols} colunas apresentam valores ausentes. "
        "Selecionamos a coluna numérica LotFrontage e a coluna categórica GarageType para detalhar a estratégia técnica adotada:"
    )

    p_num = doc.add_paragraph(style="List Bullet")
    p_num.add_run(
        f"Coluna Numérica — LotFrontage ({prep.numeric_null_count} valores nulos / 17,80% dos imóveis): "
    ).bold = True
    p_num.add_run(
        f"Foi tratada via imputação pela MEDIANA ({prep.numeric_fill_value:.1f} pés lineares). "
        "Justificativa: apagar as 259 linhas com valores nulos eliminaria quase 18% da base de treinamento, reduzindo "
        "drasticamente a capacidade de generalização do modelo e impedindo a predição caso algum dos 5 imóveis secretos "
        "não possua essa medida cadastrada. Optou-se pela mediana (69.0 ft) em vez da média (70.0 ft) porque a distribuição "
        "de LotFrontage possui assimetria positiva (outliers com terrenos superiores a 300 pés), tornando a mediana "
        "estatisticamente robusta."
    )

    p_cat = doc.add_paragraph(style="List Bullet")
    p_cat.add_run(
        f"Coluna Categórica — GarageType ({prep.categorical_null_count} valores nulos / 5,43% dos imóveis): "
    ).bold = True
    p_cat.add_run(
        "Foi tratada preenchendo os nulos com uma nova categoria explícita denominada 'None' (Sem Garagem). "
        "Justificativa: no dicionário de dados da base Ames Housing, o valor NA em GarageType (assim como em FireplaceQu, "
        "PoolQC, Alley e Fence) não representa falha de coleta, mas sim que o imóvel não possui garagem (confirmado por "
        "GarageCars = 0 e GarageArea = 0 nesses mesmos registros). Caso preenchêssemos pela moda ('Attchd' — garagem anexada), "
        "cometeríamos um erro conceitual grave, atribuindo garagem a 79 imóveis que não a possuem."
    )

    doc.add_picture(str(chart_paths["missing"]), width=Inches(6.2))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ---------------- QUESTION 3 ----------------
    doc.add_heading("3. Engenharia de Atributos (Codificação de Variáveis Categóricas)", level=1)
    q3_p = doc.add_paragraph()
    q3_p.add_run(
        "Pergunta: Transforme pelo menos duas variáveis categóricas em numéricas (usando One-Hot Encoding ou Label Encoding). "
        "Explique por que o modelo de Regressão Linear exige essa conversão."
    ).italic = True

    doc.add_paragraph(
        "Realizamos a conversão combinando as duas técnicas de acordo com a semântica de cada atributo categórico:"
    )

    p_le = doc.add_paragraph(style="List Number")
    p_le.add_run("Ordinal / Label Encoding nas variáveis ExterQual e KitchenQual: ").bold = True
    p_le.add_run(
        "Atributos que avaliam estado e qualidade possuem hierarquia natural estrita: None (0) < Po (1) < Fa (2) < TA (3) "
        "< Gd (4) < Ex (5). O mapeamento ordinal preserva a progressão numérica de qualidade em uma única coluna."
    )

    p_ohe = doc.add_paragraph(style="List Number")
    p_ohe.add_run("One-Hot Encoding nas variáveis Neighborhood, GarageType e MSZoning: ").bold = True
    p_ohe.add_run(
        "Atributos nominais sem relação de ordem foram convertidos em variáveis binárias (0 ou 1) via pd.get_dummies com "
        "drop_first=True (evitando multicolinearidade perfeita / Dummy Variable Trap)."
    )

    p_fe = doc.add_paragraph(style="List Number")
    p_fe.add_run("Engenharia de Atributos de Domínio Imobiliário (Modelo Hedônico): ").bold = True
    p_fe.add_run(
        "Adicionamos variáveis de síntese estrutural: Área Construída Total (TotalSF = TotalBsmtSF + 1stFlrSF + 2ndFlrSF), "
        "Total Ponderado de Banheiros (TotalBath), Idade do Imóvel (HouseAge), interações entre acabamento e área "
        "(Qual_x_TotalSF = OverallQual * TotalSF) e o Índice Hedônico de Preço Mediano por sq ft do Bairro (Neigh_PPSF, "
        "calculado exclusivamente sobre o conjunto de treino sem vazamento de dados). Ao todo, o espaço de features "
        f"passou de {prep.features_before_encoding} atributos para {prep.features_after_encoding} variáveis numéricas."
    )

    p_why = doc.add_paragraph()
    p_why.add_run("Por que a Regressão Linear exige essa conversão? ").bold = True
    p_why.add_run(
        "O modelo de Regressão Linear Múltipla estima o preço através da equação algébrica y = β0 + β1*x1 + β2*x2 + ... + βp*xp, "
        "resolvendo operações matriciais no espaço vetorial real (como o produto escalar X^T * X). Valores em texto (strings como "
        "'CollgCr' ou 'Gable') não admitem operações aritméticas de multiplicação por pesos (β) nem soma; portanto, é obrigatório "
        "representá-los numericamente para que o algoritmo ajuste o hiperplano ótimo."
    )

    # ---------------- QUESTION 4 ----------------
    doc.add_heading("4. Treinamento, Baseline e Interpretação de Negócios", level=1)
    q4_p = doc.add_paragraph()
    q4_p.add_run(
        "Pergunta: Separe os dados em treino e teste (ex: 80/20). Treine um modelo de Regressão Linear múltipla básico. "
        "Sem fazer ajustes avançados, qual foi o R² (R-quadrado) obtido nos dados de teste? O que esse número significa para a imobiliária?"
    ).italic = True

    r2_base = outcome.baseline_metrics["r2"]
    mae_base = outcome.baseline_metrics["mae"]
    rmse_base = outcome.baseline_metrics["rmse"]
    mape_base = outcome.baseline_metrics["mape"]

    r2_comp = outcome.competition_metrics["r2"]
    mae_comp = outcome.competition_metrics["mae"]
    rmse_comp = outcome.competition_metrics["rmse"]
    mape_comp = outcome.competition_metrics["mape"]

    doc.add_paragraph(
        f"O dataset foi particionado em 80% para treino ({outcome.train_size} imóveis) e 20% para teste ({outcome.test_size} imóveis) "
        "com semente fixa (random_state=42). Avaliamos tanto o modelo de Regressão Linear Múltipla Básico (Baseline) quanto um "
        "Pipeline Otimizado para a Competição (Ensemble de Ridge + Lasso + Gradient Boosting sobre log(SalePrice) focado em minimizar o MAE):"
    )

    table_res = doc.add_table(rows=3, cols=5)
    table_res.style = "Light Shading Accent 1"
    table_res.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Modelo Avaliado (Teste 20%)", "R² (R-Quadrado)", "MAE (US$)", "RMSE (US$)", "MAPE (%)"]
    for i, h in enumerate(headers):
        table_res.rows[0].cells[i].text = h

    r_b = table_res.rows[1].cells
    r_b[0].text = "1. Regressão Linear Múltipla (Baseline)"
    r_b[1].text = f"{r2_base:.4f} ({r2_base * 100:.2f}%)"
    r_b[2].text = f"US$ {mae_base:,.2f}"
    r_b[3].text = f"US$ {rmse_base:,.2f}"
    r_b[4].text = f"{mape_base:.2f}%"

    r_c = table_res.rows[2].cells
    r_c[0].text = "2. Pipeline Otimizado (Competição MAE)"
    r_c[1].text = f"{r2_comp:.4f} ({r2_comp * 100:.2f}%)"
    r_c[2].text = f"US$ {mae_comp:,.2f}"
    r_c[3].text = f"US$ {rmse_comp:,.2f}"
    r_c[4].text = f"{mape_comp:.2f}%"

    doc.add_paragraph()
    doc.add_picture(str(chart_paths["eval"]), width=Inches(6.2))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_meaning = doc.add_paragraph()
    p_meaning.add_run(f"Significado do R² = {r2_base:.4f} ({r2_base * 100:.2f}%) para a Imobiliária:\n").bold = True
    p_meaning.add_run(
        f"1. Poder Explicativo de Mercado: O coeficiente de determinação R² de {r2_base:.4f} significa que {r2_base * 100:.2f}% "
        "de toda a variabilidade nos preços dos imóveis é explicada pelas variáveis estruturais e de localização presentes no modelo. "
        "Para a imobiliária, isso comprova que a precificação automatizada é altamente aderente à realidade de mercado.\n"
        f"2. Risco Financeiro e Precisão Comercial: No modelo baseline, o Erro Absoluto Médio (MAE) foi de US$ {mae_base:,.2f} "
        f"({mape_base:.2f}% de desvio médio). Já utilizando o Pipeline Otimizado para a Competição, o R² atinge {r2_comp:.4f} "
        f"({r2_comp * 100:.2f}%) e o MAE cai para US$ {mae_comp:,.2f} ({mape_comp:.2f}%), garantindo estimativas precisas "
        "para evitar tanto a subprecificação (perda de comissão/patrimônio) quanto a sobreprecificação (imóvel encalhado em carteira)."
    )

    doc.save(output_docx)


def main() -> None:
    root_dir = Path(__file__).resolve().parent.parent
    csv_path = root_dir / "Trabalho" / "house-prices.csv"
    figures_dir = root_dir / "outputs" / "figures"
    models_dir = root_dir / "outputs" / "models"

    processor = HousePricesDataProcessor(csv_path)
    raw_df = processor.load_raw_data()
    X, y, prep_report = processor.fit_preprocess(raw_df)

    visualizer = HousePricesVisualizer(figures_dir)
    corr_chart = visualizer.plot_correlation_matrix(raw_df)
    scatter_chart = visualizer.plot_top3_scatter(raw_df)
    missing_chart = visualizer.plot_missing_values_treatment(raw_df)

    trainer = HousePricesModelTrainer(models_dir, random_state=42)
    outcome = trainer.train_and_evaluate(X, y, raw_df, processor)

    eval_chart = visualizer.plot_model_evaluation(
        outcome.y_test,
        outcome.y_pred_baseline,
        outcome.baseline_metrics,
        outcome.y_pred_competition,
        outcome.competition_metrics,
    )

    chart_paths = {
        "corr": corr_chart,
        "scatter": scatter_chart,
        "missing": missing_chart,
        "eval": eval_chart,
    }

    md_path = root_dir / "Trabalho" / "Relatorio_Problema1.md"
    docx_path = root_dir / "Trabalho" / "Relatorio_Problema1.docx"

    build_markdown_report(md_path, prep_report, outcome)
    build_docx_report(docx_path, prep_report, outcome, chart_paths)

    print("=== PIPELINE CONCLUÍDO COM SUCESSO ===")
    print(f"Total de linhas processadas: {prep_report.total_rows}")
    print(f"Features após encoding: {prep_report.features_after_encoding}")
    print(
        f"Baseline LinearRegression -> R²: {outcome.baseline_metrics['r2']:.4f} | "
        f"MAE: US$ {outcome.baseline_metrics['mae']:,.2f} | RMSE: US$ {outcome.baseline_metrics['rmse']:,.2f}"
    )
    print(
        f"Modelo Competição         -> R²: {outcome.competition_metrics['r2']:.4f} | "
        f"MAE: US$ {outcome.competition_metrics['mae']:,.2f} | RMSE: US$ {outcome.competition_metrics['rmse']:,.2f}"
    )
    print(f"Relatório Markdown salvo em: {md_path}")
    print(f"Relatório Word salvo em:     {docx_path}")


if __name__ == "__main__":
    main()
